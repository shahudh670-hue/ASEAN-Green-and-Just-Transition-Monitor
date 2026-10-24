"""Extract ASEAN rows from the 2026 global MPI workbooks (OPHI and UNDP) into
config/static/mpi2026.json. Run locally when OPHI publishes a new release:

    python3 scripts/mpi/extract_mpi.py <Table1.xlsx> <Table5.xlsx> <Table6.xlsx>

The workbooks are published under CC BY 4.0 at https://ophi.org.uk/publications/gMPI17-2026
"""
import json, sys, math
from pathlib import Path
import openpyxl

ASEAN = ['BRN', 'KHM', 'IDN', 'LAO', 'MYS', 'MMR', 'PHL', 'SGP', 'THA', 'TLS', 'VNM']
IND = ['nutrition', 'child_mortality', 'schooling', 'attendance', 'cooking', 'sanitation', 'water', 'electricity', 'housing', 'assets']
DIMS = ['health', 'education', 'living']


def num(v, d=4):
    if v is None or v == '' or isinstance(v, str):
        return None
    if isinstance(v, float) and math.isnan(v):
        return None
    return round(float(v), d)


def rows(wb, sheet):
    for r in wb[sheet].iter_rows(values_only=True):
        if r and len(r) > 2 and r[1] in ASEAN:
            yield r


def survey_year(y):
    return str(y).replace('-', '/') if y is not None else None


def main(t1, t5, t6, out):
    w1 = openpyxl.load_workbook(t1, read_only=True, data_only=True)
    w5 = openpyxl.load_workbook(t5, read_only=True, data_only=True)
    w6 = openpyxl.load_workbook(t6, read_only=True, data_only=True)

    national = {}
    for r in rows(w1, '1.1 National MPI Results'):
        national[r[1]] = {
            'survey': r[4], 'year': survey_year(r[5]),
            'mpi': num(r[6], 5), 'headcount': num(r[7], 3), 'intensity': num(r[8], 3),
            'vulnerable': num(r[9], 3), 'severe': num(r[10], 3), 'destitute': num(r[11], 3),
            'popSurvey': num(r[14], 1), 'pop2024': num(r[16], 1),
            'poorSurvey': num(r[17], 1), 'poor2024': num(r[19], 1),
            'nIndicators': r[20], 'missing': r[21] or '',
        }
    for r in rows(w1, '1.2 Censored Headcount Ratios'):
        national[r[1]]['censored'] = {k: num(r[7 + i], 3) for i, k in enumerate(IND)}
    for r in rows(w1, "1.3 Contribut'n of deprivations"):
        national[r[1]]['contribDims'] = {k: num(r[7 + i], 2) for i, k in enumerate(DIMS)}
        national[r[1]]['contrib'] = {k: num(r[10 + i], 2) for i, k in enumerate(IND)}
    for r in rows(w1, '1.6 Raw Headcount Ratios'):
        national[r[1]]['raw'] = {k: num(r[7 + i], 3) for i, k in enumerate(IND)}

    regions = {}
    for r in rows(w5, '5.1 MPI Results'):
        key = (r[1], r[6].strip())
        regions[key] = {
            'iso': r[1], 'name': r[6].strip(), 'year': survey_year(r[5]),
            'mpi': num(r[8], 5), 'headcount': num(r[9], 3), 'intensity': num(r[10], 3),
            'vulnerable': num(r[11], 3), 'severe': num(r[12], 3),
            'popShare': num(r[16], 3), 'pop2024': num(r[17], 1), 'poor2024': num(r[18], 1),
        }
    for sheet, field, start in [('5.2 Censored Headcount Ratios', 'censored', 9), ('5.5 Raw Headcount Ratios', 'raw', 9)]:
        for r in rows(w5, sheet):
            regions[(r[1], r[6].strip())][field] = {k: num(r[start + i], 3) for i, k in enumerate(IND)}
    for r in rows(w5, "5.3 Contribut'n of deprivations"):
        g = regions[(r[1], r[6].strip())]
        g['contribDims'] = {k: num(r[9 + i], 2) for i, k in enumerate(DIMS)}
        g['contrib'] = {k: num(r[12 + i], 2) for i, k in enumerate(IND)}

    def block(r, c):
        return {'t0': num(r[c], 5), 't1': num(r[c + 1], 5), 'abs': num(r[c + 2], 5), 'rel': num(r[c + 3], 3), 'sig': (r[c + 5] or '').strip()}

    changes = []
    for r in rows(w6, '6.1 Harmonised MPI'):
        ch = {
            'iso': r[1], 'survey0': r[4], 'year0': survey_year(r[5]), 'survey1': r[6], 'year1': survey_year(r[7]),
            'mpi': block(r, 8), 'headcount': block(r, 14), 'intensity': block(r, 20),
            'vulnerable': block(r, 26), 'severe': block(r, 32),
            'poor0': num(r[42], 1), 'poor1': num(r[43], 1),
            'censored': {k: block(r, 48 + 6 * i) for i, k in enumerate(IND)},
            'nIndicators': r[108], 'missing': r[109] or '',
        }
        changes.append(ch)

    # Identity checks: MPI = H x A, and indicator contributions sum to 100.
    for iso, n in national.items():
        assert abs(n['mpi'] - n['headcount'] * n['intensity'] / 1e4) < 2e-3, iso
        assert abs(sum(v for v in n['contrib'].values() if v is not None) - 100) < 0.5, iso
    for g in regions.values():
        assert abs(g['mpi'] - g['headcount'] * g['intensity'] / 1e4) < 2e-3, g['name']

    doc = {
        'title': '2026 global Multidimensional Poverty Index, ASEAN members',
        'publisher': 'Oxford Poverty and Human Development Initiative (OPHI), University of Oxford, and UNDP Human Development Report Office',
        'licence': 'CC BY 4.0',
        'pageUrl': 'https://ophi.org.uk/publications/gMPI17-2026',
        'hdroUrl': 'https://hdr.undp.org/content/2026-global-multidimensional-poverty-index',
        'citations': {
            'national': "Alkire, S., Mishra, R. and Suppa, N. (2026). 'The global Multidimensional Poverty Index (MPI) 2026 country results and methodological note', OPHI MPI Methodological Note 64, OPHI, University of Oxford. Table 1.",
            'regions': "Alkire, S., Mishra, R. and Suppa, N. (2026). 'The global Multidimensional Poverty Index (MPI) 2026: Disaggregation results and methodological note', OPHI MPI Methodological Note 65, OPHI, University of Oxford. Table 5.",
            'changes': "Alkire, S., Mishra, R. and Suppa, N. (2026). 'A methodological note on the global Multidimensional Poverty Index (MPI) 2026: Changes over time results for 89 countries', OPHI MPI Methodological Note 66, OPHI, University of Oxford. Table 6.1.",
        },
        'units': 'mpi 0 to 1; headcount, intensity, vulnerable, severe, destitute, censored and raw in % of population; contributions in % of the MPI; population and poor in thousands. poor2024 applies the survey headcount to the 2024 population.',
        'indicators': IND,
        'notCovered': [i for i in ASEAN if i not in national],
        'national': national,
        'regions': sorted(regions.values(), key=lambda g: (g['iso'], g['name'])),
        'changes': changes,
    }
    Path(out).write_text(json.dumps(doc, ensure_ascii=False, indent=1))
    print(f"national {len(national)}, regions {len(regions)}, changes {len(changes)} -> {out}")


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[2]
    main(*sys.argv[1:4], out=root / 'config/static/mpi2026.json')
